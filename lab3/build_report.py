"""Build the laboratory report after app screenshots and evaluation are ready."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).parent
ASSETS = ROOT / "report_assets"
REPORT = ROOT / "Отчёт лабораторная работа 3 вариант 18.docx"
EVALUATION = json.loads((ROOT / "evaluation.json").read_text(encoding="utf-8"))
MANIFEST = json.loads((ROOT / "corpus" / "manifest.json").read_text(encoding="utf-8"))


def make_diagram():
    path = ASSETS / "algorithm.png"
    image = Image.new("RGB", (1300, 960), "white")
    draw = ImageDraw.Draw(image)
    font_path = "C:/Windows/Fonts/arial.ttf"
    font = ImageFont.truetype(font_path, 29)
    small = ImageFont.truetype(font_path, 24)
    blocks = [
        ("Текст + язык + корпус", "Входные данные"),
        ("Абзацы → предложения → слова", "Разбор документа"),
        ("Частоты tf и df → модифицированная TF-IDF", "Веса терминов"),
        ("Σ tf(t, Si) × w(t, D) × Posd × Posp", "Sentence extraction"),
        ("Топ-10 в исходном порядке + топ-12 слов", "Результат"),
    ]
    for i, (main, subtitle) in enumerate(blocks):
        y = 35 + i * 185
        draw.rounded_rectangle((90, y, 1210, y + 125), radius=22,
                               fill="#edf3ff", outline="#5076c3", width=3)
        draw.text((130, y + 18), subtitle, font=small, fill="#2d4a7a")
        draw.text((130, y + 57), main, font=font, fill="#17243b")
        if i < len(blocks) - 1:
            draw.line((650, y + 126, 650, y + 176), fill="#5076c3", width=5)
            draw.polygon([(638, y + 164), (662, y + 164), (650, y + 180)], fill="#5076c3")
    image.save(path)
    return path


def set_cell(cell, text):
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(9)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, title in zip(table.rows[0].cells, headers):
        set_cell(cell, title)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell(cell, value)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph()


def caption(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.style = "Caption"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True


def figure(doc, path, text, width=14.5):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(path), width=Cm(width))
    caption(doc, text)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 19), ("Heading 1", 14), ("Heading 2", 11.5)):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor(20, 38, 68)

    doc.add_paragraph("Отчёт по лабораторной работе 3", "Title")
    doc.add_paragraph("Автоматическое реферирование документов  вариант 18", "Subtitle")
    doc.add_paragraph("Цель работы: реализовать и проверить извлечение значимых предложений и ключевых слов для английских и испанских текстов по медицине и критике изобразительного искусства.")
    doc.add_paragraph("Результат: создано Flask-приложение с методом Sentence extraction по формулам методички и дополнительным методом латентно семантического анализа (LSA), выбранным вместо OSTIS согласно заданию. Выход содержит ссылку на источник, классический реферат и список ключевых слов.")
    doc.add_heading("1 Постановка задачи и соответствие требованиям", 1)
    add_table(doc, ["Требование", "Реализация и проверка"], [
        ("Вариант 18", "2 языка × 2 области: английский, испанский; медицина, критика искусства."),
        ("Тексты одинакового размера", "4 файла UTF-8 по 30 предложений; объём 375–406 слов (разброс 8%)."),
        ("Sentence extraction", "Модифицированная TF-IDF, положение в документе и абзаце, выбор 10 предложений в исходном порядке."),
        ("Замена OSTIS", "LSA на матрице предложение–термин с сингулярным разложением."),
        ("Два раздела результата", "Классический реферат и 12 ключевых слов TF-IDF."),
        ("Хранение больших коллекций", "PostgreSQL: исходные документы, рефераты и ключевые слова; связь через внешние ключи."),
        ("Источник, сохранение, печать", "Активная ссылка на исходный текст из PostgreSQL; экспорт .txt; кнопка печати браузера."),
        ("Простой интерфейс и help", "Bootstrap, форма выбора, загрузка своего .txt, отдельная страница справки."),
    ], [5.5, 10])

    doc.add_heading("2 Тестовая коллекция", 1)
    doc.add_paragraph("Коллекция находится в папке corpus. Тексты составлены специально для проверки, поэтому оценка не претендует на результат для реальных научных статей или опубликованных рецензий. Во всех файлах шесть абзацев и 30 предложений. Размер сравним по числу слов; около десяти страниц в методичке приведено как пример, а не обязательный нижний предел.")
    rows = []
    for item in MANIFEST:
        text = (ROOT / "corpus" / item["file"]).read_text(encoding="utf-8")
        rows.append((item["file"], item["language"].upper(), item["domain"], len(text.split()), 30))
    add_table(doc, ["Файл", "Язык", "Область", "Слов", "Предл."], rows)
    doc.add_paragraph("Медицинские тексты рассматривают сон и сердечно-сосудистое здоровье; художественные — композицию, свет и фактуру условной пейзажной картины. Пары языков тематически сопоставимы, но сформулированы независимо.")

    doc.add_heading("3 Структура системы и данные", 1)
    doc.add_paragraph("app.py обрабатывает HTTP-запросы и выводит шаблоны Jinja2. summarizer.py делит текст на абзацы, предложения и термины, вычисляет веса и формирует результат. corpus/manifest.json хранит имя, язык, область и название каждого демонстрационного примера. Файлы в corpus/*.txt содержат исходную тестовую коллекцию. Загруженные пользователем документы, рефераты и ключевые слова хранятся в PostgreSQL через модуль database.py.")
    doc.add_paragraph("В базе используются три связанные таблицы. documents хранит исходный файл: имя, язык, предметную область, полный текст и дату загрузки. summaries хранит результат для документа и метода: число предложений, классический реферат и список ключевых слов в JSONB. summary_keywords содержит нормализованный список ключевых слов с весом TF-IDF и позицией. При удалении документа связанные рефераты и ключевые слова удаляются каскадно. Пара (document_id, method, sentence_count) уникальна, поэтому результат одного метода с одинаковыми параметрами не дублируется.")
    doc.add_paragraph("В памяти предложение представлено структурой Sentence: текст, смещение в документе, смещение и длина абзаца, порядковый номер. Частоты tf и df хранятся в Counter, веса терминов — в словаре. Выход функции summarize — словарь с массивами предложений и ключевых слов, языком, методом и общим числом предложений. HTML и .txt формируются из этого результата.")
    doc.add_heading("4 Алгоритмы построения реферата", 1)
    doc.add_heading("4.1 Sentence extraction", 2)
    doc.add_paragraph("Текст разбивается по абзацам и знакам конца предложения. Числа и стоп-слова исключаются. Указание методички не учитывать латинские слова относится к иной языковой конфигурации: для варианта 18 оно неприменимо, поскольку оба языка используют латиницу. Поэтому английские и испанские слова сохраняются.")
    doc.add_paragraph("Вес термина: w(t,D) = 0,5 × (1 + tf(t,D)/tfmax(D)) × ln(|DB|/df(t)). Здесь DB — документы того же языка; df(t) — число содержащих термин документов. Вес предложения: Score(Si) = Σ tf(t,Si) × w(t,D). Итоговый вес: Score(Si) × [1 − BD(Si)/|D|] × [1 − BP(Si)/|P|]. Смещения измеряются в символах от начала текста и абзаца.")
    doc.add_paragraph("Выбираются 10 предложений с максимальным итоговым весом, затем они сортируются по исходным номерам. Отдельно берутся 12 терминов с наибольшим положительным весом w(t,D).")
    figure(doc, make_diagram(), "Рисунок 1  Последовательность обработки для Sentence extraction", 14.5)
    doc.add_heading("4.2 Латентно семантический анализ", 2)
    doc.add_paragraph("Дополнительный метод строит TF-IDF-матрицу предложений, сокращает её до не более чем трёх компонент с помощью TruncatedSVD и оценивает косинусную близость каждого предложения к центру текста в полученном пространстве. Он выбирает 10 предложений, сохраняя их исходный порядок. Раздел ключевых слов остаётся общим для обоих методов и строится по документной TF-IDF.")

    doc.add_heading("5 Проверка и оценка", 1)
    doc.add_paragraph("Проверены главная страница, справка, все 4 исходных файла, оба метода для каждого файла, ссылка на источник, загрузка пользовательского файла, сохранение и наличие кнопки печати. Для загрузки в тестах подменяются функции PostgreSQL, что изолирует обработчик формы от локальных учётных данных БД. Автоматический набор содержит 5 тестов unittest; все прошли. Для оценки содержания вручную отмечены по 8 опорных предложений в каждом тексте. Precision@10 = число выбранных опорных / 10, Recall@8 = число выбранных опорных / 8. Время — среднее 30 запусков на локальной машине, без HTTP и отрисовки браузера.")
    rows = []
    for row in EVALUATION:
        rows.append((row["file"].replace(".txt", ""), "SE" if row["method"] == "sentence" else "LSA",
                     f'{row["precision_at_10"]:.2f}', f'{row["recall_at_8"]:.2f}',
                     f'{row["mean_ms"]:.2f}'))
    add_table(doc, ["Документ", "Метод", "P@10", "R@8", "мс"], rows)
    doc.add_paragraph("На этой коллекции Sentence extraction даёт P@10 = 0,30 для всех четырёх файлов. LSA даёт 0,10–0,40 в зависимости от текста. Оценка ограничена маленькой авторской коллекцией и субъективной ручной разметкой. Позиционный множитель Sentence extraction заметно предпочитает начало текста, из-за чего поздние важные выводы попадают в реферат реже. LSA занимает больше времени, но остаётся быстрым на этих коротких текстах.")

    doc.add_heading("6 Интерфейс и результаты", 1)
    figure(doc, ASSETS / "home.png", "Рисунок 2  Выбор документа и метода на главной странице", 13.2)
    figure(doc, ASSETS / "result_sentence.png", "Рисунок 3  Реферат медицинского текста методом Sentence extraction", 11.7)
    figure(doc, ASSETS / "result_lsa.png", "Рисунок 4  Реферат испанского текста об искусстве методом LSA", 12.5)
    figure(doc, ASSETS / "help.png", "Рисунок 5  Страница справки", 13.2)

    doc.add_heading("7 Применённые компоненты", 1)
    doc.add_paragraph("Flask предоставляет маршруты и обработку форм; Jinja2 формирует страницы. Bootstrap 5.3.3 отвечает за сетку, формы и кнопки; его стили подключены через CDN. psycopg 3 выполняет параметризованные запросы к PostgreSQL и передаёт результаты в JSONB. scikit-learn предоставляет TfidfVectorizer и TruncatedSVD для LSA, NumPy — операции с векторами. Стандартные pathlib, re, collections и math используются для файлов, сегментации и формул. python-docx применён только для подготовки этого отчёта, не является зависимостью приложения.")
    doc.add_heading("8 Выводы", 1)
    doc.add_paragraph("Система выполняет требования варианта 18: поддерживает два языка и две предметные области, строит два вида выходной информации, открывает источник, сохраняет и печатает результат, содержит справку. PostgreSQL обеспечивает хранение загружаемых документов и результатов и позволяет масштабировать систему на большую коллекцию без хранения файлов в папке приложения. Метод Sentence extraction реализован по формулам задания; LSA даёт разрешённую замену OSTIS. Проверка выявила ограничение позиционного ранжирования и зависимость LSA от тематики. Для дальнейшего применения полезны более крупный корпус, полноценная лемматизация английского и испанского языков и оценка на независимых экспертных рефератах.")
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    build()
